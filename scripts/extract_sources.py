"""Read-only extraction of the four user-supplied reference files."""
import json, re, hashlib
from pathlib import Path
import openpyxl
from docx import Document
from lxml import html

base = Path('/Users/kadirerbakar/Downloads')
out = Path(__file__).resolve().parents[1] / 'data'
files = [base/'oyuncak gtip listesi.xlsx', next(p for p in base.glob('oyuncak*.xlsx') if 'vergisi' in p.name), base/'ithalatta_gozetim_uygulanmasina_iliskin_teblig_2019-06.docx', base/'2026-10_ugd_oyuncaklarin_ithalat_denetimi_tebligi.htm']
clean = lambda x: re.sub(r'\s+', ' ', str(x or '')).strip()
codes = {}
path = []
for rownum, row in enumerate(openpyxl.load_workbook(files[0], data_only=True).active.values, 1):
    code, desc = clean(row[0]), clean(row[1])
    if re.fullmatch(r'9503\.\d{2}\.\d{2}\.\d{2}\.\d{2}', code):
        label = desc.lstrip('- ')
        context = ' / '.join(path)
        codes[code] = {'code':code, 'label':label, 'context':context, 'sourceRow':rownum}
    elif not code and desc.startswith('-'):
        depth = len(re.match(r'^[- ]+', desc).group().replace(' ', ''))
        path = path[:depth-1] + [desc.lstrip('- ')]
groups = ['PREF','KTR','BAE','EAGU','OTDU','GYU','DU']
kind = 'customs'
for rownum, row in enumerate(openpyxl.load_workbook(files[1], data_only=True).active.values, 1):
    code = clean(row[0])
    if code.startswith('İlave'): kind = 'additional'
    if code in codes:
        codes[code][kind] = dict(zip(groups, row[3:10]))
        codes[code][kind+'SourceRow'] = rownum
surveillance = []
for rownum,row in enumerate(Document(files[2]).tables[0].rows[1:],2):
    code, desc, rate = [clean(c.text) for c in row.cells]
    surveillance.append({'code':code, 'label':desc, 'usdPerGrossTonne':int(rate.replace('.','')), 'sourceTable':1, 'sourceRow':rownum})
root=html.parse(str(files[3]))
tareks={}
for tr in root.xpath('//tr'):
    cells=[clean(t.text_content()) for t in tr.xpath('./td')]
    if len(cells)>=3 and re.fullmatch(r'\d{4}\.\d{2}\.\d{2}\.\d{2}\.\d{2}',cells[1]): tareks[cells[1]]=cells[2]
corrections=json.loads((out/'source-corrections.json').read_text())
issues=[]
for code, product in codes.items():
    exact=[r for r in surveillance if r['code']==code]
    product['surveillanceRate']=exact[0]['usdPerGrossTonne'] if len(exact)==1 else None
    product['surveillanceStatus']='multiple' if len(exact)>1 else ('listed' if exact else 'not-listed')
    product['surveillanceSourceRows']=[r['sourceRow'] for r in exact]
    for correction in corrections:
        if code!=correction['targetCode']: continue
        source_rows=[r for r in surveillance if r['code']==correction['sourceCode']]
        assert sorted(r['usdPerGrossTonne'] for r in source_rows)==sorted(v['usdPerGrossTonne'] for v in correction['variants']), 'Correction source changed; review required'
        product['surveillanceStatus']='variant-required'
        product['surveillanceVariants']=correction['variants']
        product['surveillanceCorrection']=correction
        product['surveillanceSourceRows']=[r['sourceRow'] for r in source_rows]
        issues.append({'code':code,'status':'corrected-with-provenance','reason':correction['note']})
    product['tareksDescription']=tareks.get(code)
    assert 'customs' in product and 'additional' in product
payload={'version':'2026-10-01-source-audit-v2','effectiveFrom':'2026-08-10','snapshotDate':'2026-10-01','vatRate':20,'products':list(codes.values()),'surveillanceRows':surveillance,'sources':[{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files], 'issues':issues}
(out/'catalog.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2))
print(f'{len(codes)} GTİP, {len(surveillance)} gözetim satırı, {len(tareks)} denetim satırı aktarıldı.')
